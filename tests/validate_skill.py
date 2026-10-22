"""Structure and rule checks for the english-classroom plugin. Stdlib only."""
import json
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SKILL_DIR = ROOT / "skills" / "english-classroom"
SKILL_MD = SKILL_DIR / "SKILL.md"
REFS = SKILL_DIR / "references"
EM_DASH = "\u2014"


def read(path):
    return Path(path).read_text(encoding="utf-8")


def frontmatter(text):
    match = re.match(r"---\n(.*?)\n---\n", text, re.S)
    if not match:
        return {}
    fields = {}
    for line in match.group(1).splitlines():
        if ":" in line:
            key, value = line.split(":", 1)
            fields[key.strip()] = value.strip()
    return fields


def non_heading_lines(text):
    body = re.sub(r"\A---\n.*?\n---\n", "", text, flags=re.S)
    return [ln for ln in body.splitlines() if not ln.lstrip().startswith("#")]


def em_dash_lines(path):
    return [ln for ln in non_heading_lines(read(path)) if EM_DASH in ln]


class TestScaffold(unittest.TestCase):
    def test_required_files_exist(self):
        required = [
            ".claude-plugin/plugin.json",
            ".claude-plugin/marketplace.json",
            ".gitignore",
            "LICENSE",
            "skills/english-classroom/scripts/md_to_docx.py",
            "skills/english-classroom/references/cefr-levels.md",
            "skills/english-classroom/references/lesson-frameworks.md",
            "skills/english-classroom/references/activity-bank.md",
            "skills/english-classroom/references/teacher-activities.md",
        ]
        for rel in required:
            self.assertTrue((ROOT / rel).is_file(), rel)

    def test_plugin_json(self):
        data = json.loads(read(ROOT / ".claude-plugin/plugin.json"))
        self.assertEqual(data["name"], "english-classroom")
        self.assertEqual(data["version"], "1.1.0")
        self.assertEqual(data["license"], "MIT")
        self.assertIn("ratherlegit/english-classroom", data["repository"])

    def test_marketplace_json_matches_plugin(self):
        plugin = json.loads(read(ROOT / ".claude-plugin/plugin.json"))
        market = json.loads(read(ROOT / ".claude-plugin/marketplace.json"))
        self.assertEqual(market["name"], "english-classroom")
        self.assertEqual(market["version"], plugin["version"])
        self.assertEqual(market["plugins"][0]["name"], "english-classroom")
        self.assertEqual(market["plugins"][0]["source"], "./")

    def test_gitignore(self):
        self.assertIn(".DS_Store", read(ROOT / ".gitignore"))

    def test_converter_identical_to_english_tutoring_copy(self):
        src = Path.home() / ".claude/skills/english-tutoring/skills/english-tutoring/scripts/md_to_docx.py"
        if not src.is_file():
            self.skipTest("english-tutoring not installed locally")
        self.assertEqual(
            (SKILL_DIR / "scripts/md_to_docx.py").read_bytes(), src.read_bytes()
        )


class TestExport(unittest.TestCase):
    def test_converter_builds_docx_with_heading_and_table(self):
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx not installed")
        md = (
            "# Test Lesson\n\n## Warm Up Conversation\n\n- One\n- Two\n\n"
            "| Lesson # | Lesson Name | Grouping | Status |\n|---|---|---|---|\n"
            "| 1 | Intro | Mixed pairs | Not yet delivered |\n"
        )
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "plan.md"
            out = Path(tmp) / "plan.docx"
            src.write_text(md, encoding="utf-8")
            subprocess.run(
                [sys.executable, str(SKILL_DIR / "scripts/md_to_docx.py"), str(src), str(out)],
                check=True,
            )
            self.assertTrue(out.is_file())
            document = docx.Document(str(out))
            self.assertEqual(len(document.tables), 1)
            texts = [p.text for p in document.paragraphs]
            self.assertIn("Test Lesson", texts)


class TestNewReferences(unittest.TestCase):
    def test_class_interview_content(self):
        text = read(REFS / "class-interview.md")
        for phrase in [
            "Number of students",
            "Backgrounds",
            "mono-lingual",
            "mixed",
            "Light mode",
            "Roster",
            "L1 interference cheat sheet",
            "age-appropriate",
            "can do",
            "headcount",
            "first name",
            "last initial",
            "diagnoses",
        ]:
            self.assertIn(phrase, text, phrase)

    def test_differentiation_content(self):
        text = read(REFS / "differentiation.md")
        for phrase in [
            "Support",
            "Core",
            "Stretch",
            "two CEFR bands",
            "parallel",
            "Small",
            "Medium",
            "Large",
            "mixed-level pairs",
            "early finishers",
            "single level",
            "never a shorter",
            "bimodal",
            "headcount",
            "lower of the two",
            "Quick Level Check",
        ]:
            self.assertIn(phrase, text, phrase)

    def test_no_em_dashes_in_new_references(self):
        for name in ["class-interview.md", "differentiation.md"]:
            self.assertEqual(em_dash_lines(REFS / name), [], name)


def group_section(name):
    text = read(REFS / name)
    marker = "## Group classroom notes"
    return text[text.index(marker):]


class TestGroupNotes(unittest.TestCase):
    def test_group_notes_appended(self):
        for name in ["teacher-activities.md", "activity-bank.md", "lesson-frameworks.md"]:
            self.assertIn("## Group classroom notes", read(REFS / name), name)

    def test_teacher_activities_group_rules(self):
        text = group_section("teacher-activities.md")
        for phrase in ["Support", "Stretch", "rotate", "Press conference", "Found article"]:
            self.assertIn(phrase, text, phrase)

    def test_group_sections_no_em_dashes(self):
        for name in ["teacher-activities.md", "activity-bank.md", "lesson-frameworks.md"]:
            self.assertNotIn(EM_DASH, group_section(name), name)


class TestSkillMd(unittest.TestCase):
    def setUp(self):
        self.text = read(SKILL_MD)

    def test_frontmatter(self):
        fm = frontmatter(self.text)
        self.assertEqual(fm.get("name"), "english-classroom")
        desc = fm.get("description", "")
        self.assertIn("english-tutoring", desc)
        self.assertIn("class", desc)
        self.assertLess(len(desc), 1024)

    def test_steps_present_in_order(self):
        headings = [
            "## Step 0: Check for a returning class",
            "## Step 1: Interview the teacher about the class",
            "## Step 2: Design the class learning path",
            "## Step 3: Pick a framework",
            "## Step 4: Build the plan",
            "## Step 5: Differentiation and common pitfalls",
            "## Output format",
            "## Step 6: Export as a .docx (optional)",
        ]
        positions = [self.text.find(h) for h in headings]
        self.assertTrue(all(p >= 0 for p in positions), positions)
        self.assertEqual(positions, sorted(positions))

    def test_required_rules(self):
        for phrase in [
            "~/eng-teaching-records/classes/",
            "median",
            "Support",
            "Core",
            "Stretch",
            "Grouping",
            "two CEFR bands",
            "parallel",
            "never a shorter",
            "exit task",
            "bimodal",
            "headcount",
            "new term",
            "path is complete",
            "Quick Level Check",
            "last initial",
            "estimated",
            "Never overwrite a record",
            "Previous Paths",
            "whole roster",
            "never inferred",
            "culture-sharing prompt",
            "Differentiation",
            "Review of Prior Content",
            "Activity Log",
            "American by default",
            "no student-facing text may mention the teacher",
            "Exports are black text only",
            "never lettered or numbered",
            "first lesson",
            "Never save exports inside",
        ]:
            self.assertIn(phrase, self.text, phrase)

    def test_review_focus_rules(self):
        for phrase in [
            "class of one",
            "already exists",
            "age-appropriate",
            "single level",
            "can do",
            "cannot generate a Word file",
            "diagnoses",
        ]:
            self.assertIn(phrase, self.text, phrase)

    def test_referenced_files_exist(self):
        for rel in re.findall(r"`((?:references|scripts)/[\w./-]+)`", self.text):
            self.assertTrue((SKILL_DIR / rel).is_file(), rel)

    def test_no_em_dashes(self):
        self.assertEqual(em_dash_lines(SKILL_MD), [])


class TestDocs(unittest.TestCase):
    def test_readme(self):
        text = read(ROOT / "README.md")
        self.assertIn("ratherlegit/english-classroom", text)
        self.assertIn("english-tutoring", text)
        self.assertIn("claude plugin install english-classroom@english-classroom", text)
        self.assertIn("`/english-classroom`", text)
        self.assertEqual(em_dash_lines(ROOT / "README.md"), [])

    def test_agents(self):
        text = read(ROOT / "AGENTS.md")
        self.assertIn("skills/english-classroom/SKILL.md", text)
        self.assertIn("~/eng-teaching-records/classes/", text)
        self.assertEqual(em_dash_lines(ROOT / "AGENTS.md"), [])


class TestPublishHygiene(unittest.TestCase):
    def test_no_local_paths_or_log_em_dashes(self):
        files = [ROOT / "README.md", ROOT / "AGENTS.md", SKILL_MD,
                 REFS / "class-interview.md", REFS / "differentiation.md",
                 ROOT / "tests" / "manual-test-log.md"]
        files += [p for p in (ROOT / "docs").rglob("*") if p.is_file()]
        for path in files:
            self.assertNotIn("/Users/", read(path), str(path))
        self.assertNotIn(EM_DASH, read(ROOT / "tests" / "manual-test-log.md"))


if __name__ == "__main__":
    unittest.main()
